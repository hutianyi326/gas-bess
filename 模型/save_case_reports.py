"""归档两组单年已收敛结果并生成独立报告，不重新求解。"""
from pathlib import Path
from datetime import datetime
import csv
import hashlib
import json
import re
import shutil
import numpy as np
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / '优化缓存/储+氢/分时电价_2026_年用氢3万吨_无绿电约束_v2/单年筛选'
OUT = ROOT / '报告/储能制氢单年定容测算_20261015'
OUT.mkdir(exist_ok=True)
REF = ROOT / '报告/轮台绿电制氢项目定容测算报告.docx'
NAMES = ['光伏', '需量容量费', '电解槽', '罐体', '泵体', '储能2h', '储能4h']


def native(x):
    if isinstance(x, dict): return {k: native(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)): return [native(v) for v in x]
    if isinstance(x, np.ndarray): return x.tolist()
    if isinstance(x, np.generic): return x.item()
    return x


def archive(label, factor, source, tag):
    dest = OUT / label
    dest.mkdir(exist_ok=True)
    raw = dest / '原始结果'; raw.mkdir(exist_ok=True)
    for p in source.iterdir():
        if p.is_file() and (p.name.startswith(tag) or p.name == '定容运行日志.txt'):
            shutil.copy2(p, raw / p.name)
    text = (source / '定容运行日志.txt').read_text(encoding='utf-8-sig')
    capacity = source / (tag + '_容量.npy')
    if capacity.exists():
        d = np.load(capacity, allow_pickle=True).item()
        x, total, coeff, gap = d['x'], d['UB'], d['capex'], d['gap']
        parts = dict(zip(NAMES, coeff * x / 1e4))
        parts['下网电量电费'] = (total - coeff @ x) / 1e4
    else:
        d = np.load(source / (tag + '_割平面.npy'), allow_pickle=True).item()
        x, total, coeff = d['best_x'], d['UB'], d['capex']
        found = re.findall(r'间隙 ([\deE.+-]+)，信赖域', text)
        gap = float(found[-1])
        parts = dict(zip(NAMES, coeff * x / 1e4))
        parts['下网电量电费'] = (total - coeff @ x) / 1e4
    # 从与最佳容量完全对应的子问题割平面验证年度成本，不以最后一轮候选代替最优点。
    checkpoint = np.load(source / (tag + '_割平面.npy'), allow_pickle=True).item()
    matches = [c for c in checkpoint['cuts'] if np.allclose(c[1], x, rtol=1e-10, atol=1e-7)]
    assert matches and abs(float(matches[-1][2]) - (total - coeff @ x)) < 1.0
    initial = np.array([x[0]*2600*factor, 0, x[2]*12000*factor,
                        x[3]/11.1259*3000*factor, x[4]/11.1259*30000*factor,
                        x[5]*780, x[6]*700])/1e4
    rows = re.findall(r'第\s*(\d+)轮.*?间隙 ([\deE.+-]+).*?罚项 ([\d.]+) GWh', text)
    s = dict(case=label, equipment_cost_factor=factor, sizing_years=1,
             maintenance=False, annual_hydrogen_t=30000, capacity=x,
             gap=gap, gap_source='容量文件' if capacity.exists() else '最后一轮日志精度',
             iterations=int(rows[-1][0]), cost_wanyuan_year=total/1e4,
             cost_parts_wanyuan_year=parts, initial_investment_wanyuan=initial,
             initial_investment_total_wanyuan=initial.sum(),
             unit_cost_yuan_kg=total/30000000,
             main_optimization_converged=gap<1e-4,
             full_process_completed=capacity.exists() and '总用时' in text,
             final_iteration_slack_GWh=float(rows[-1][2]),
             source_directory=str(source), source_tag=tag,
             caveat='设备投资含建安范围未明确；缺少水耗等费用，不代表完整平准化制氢成本')
    # 为中止算例导出明确标记的归档数据，不伪装为原程序完整输出。
    (dest / '测算结果汇总.json').write_text(json.dumps(native(s), ensure_ascii=False, indent=2), encoding='utf-8')
    with (dest / '设备配置及投资.csv').open('w', encoding='utf-8-sig', newline='') as f:
        w=csv.writer(f);w.writerow(['项目','模型容量','单位','初始投资万元','年化费用万元每年'])
        units=['kWp','kW','Nm³/h','Nm³','Nm³/h','kWh','kWh']
        for i,n in enumerate(NAMES):w.writerow([n,float(x[i]),units[i],float(initial[i]),float(coeff[i]*x[i]/1e4)])
    with (dest / '年均成本.csv').open('w', encoding='utf-8-sig', newline='') as f:
        w=csv.writer(f);w.writerow(['成本项','万元每年']);w.writerows(parts.items());w.writerow(['合计',total/1e4])
    if not capacity.exists():
        np.save(dest / '最佳收敛检查点_归档.npy', dict(x=x, UB=total, gap=gap, capex=coeff, parts=parts,
                source='原割平面检查点；后处理已人为终止；非原程序完整容量输出'), allow_pickle=True)
    inputs=dest/'输入与代码快照'; inputs.mkdir(exist_ok=True)
    shutil.copy2(ROOT/'处理后输入/储+氢_分时电价_2026.csv',inputs/'储+氢_分时电价_2026.csv')
    for name in ['lp_multi.py','bd10y.py','paths.py','requirements.txt']:
        shutil.copy2(ROOT/'模型'/name,inputs/name)
    reference_text = '\n'.join(p.text for p in Document(REF).paragraphs if p.text.strip())
    (dest/'参考报告正文.txt').write_text(reference_text,encoding='utf-8')
    return dest,s


def font(run, size=11, bold=False, face='宋体'):
    run.font.name=face;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),face)


def para(d,text,bold=False):
    p=d.add_paragraph();p.paragraph_format.space_after=Pt(7);p.paragraph_format.line_spacing=1.25
    font(p.add_run(text),11,bold)
    return p


def heading(d,text):
    p=d.add_paragraph(style='Heading 1');p.paragraph_format.space_before=Pt(10);p.paragraph_format.space_after=Pt(7)
    if text.startswith(('三 ', '五 ', '七 ')):
        p.paragraph_format.page_break_before=True
    font(p.add_run(text),14,True,'黑体');return p


def table(d,headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,h in zip(t.rows[0].cells,headers):c.text=h
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
            pr.append(b);mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF2');pr.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.1
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:font(r,10.5,ri==0)
    d.add_paragraph().paragraph_format.space_after=Pt(2)
    return t


def build(dest,s,base):
    factor=s['equipment_cost_factor'];x=np.array(s['capacity']);p=s['cost_parts_wanyuan_year'];ini=s['initial_investment_wanyuan']
    title='储能制氢项目单年定容测算报告'
    subtitle='基准造价方案' if factor==1 else '非储能设备造价降低百分之四十方案'
    d=Document();sec=d.sections[0]
    # 参考报告使用 A4；保留中文技术报告页面规格，正文改为适合本次独立报告的密度。
    sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.3);sec.bottom_margin=Cm(2.1);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.5)
    for style in ['Normal','Title','Heading 1','Heading 2']:
        st=d.styles[style];st.font.name='宋体';st.font.color.rgb=RGBColor(0,0,0)
    for border in list(d.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    hp=sec.header.paragraphs[0];font(hp.add_run('储能制氢项目  单年无检修工况'),9)
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=fp.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
    h=d.add_paragraph(style='Title');font(h.add_run(title),21,True,'黑体')
    para(d,subtitle,True);para(d,'编制日期  2026年10月15日')
    heading(d,'一 测算目的与结论')
    para(d,f'本报告测算无光伏且不含炼化用电负荷的储能制氢系统，按一年8760小时、全年供氢30000吨、无检修工况联合优化设备容量与逐时运行。{subtitle}下，2小时与4小时储能最优容量均为零；模型年均成本为{s["cost_wanyuan_year"]:,.2f}万元。')
    para(d,f'定容经过{s["iterations"]}轮迭代，最优性间隙为{s["gap"]*100:.5f}%，低于0.01%的停止阈值。该精度对应当前模型，不代表工程定型精度。')
    heading(d,'二 最优设备配置')
    table(d,['项目','最优配置','初始投资 万元'],[
        ['光伏','0 MWp','0.00'],['申报需量',f'{x[1]/1000:.2f} MW','不计资产投资'],
        ['电解槽',f'{x[2]:,.2f} Nm³/h',f'{ini[2]:,.2f}'],
        ['可用储氢容量',f'{x[3]/11.1259/1000:.2f} 吨',f'{ini[3]:,.2f}'],
        ['增压机',f'{x[4]/11.1259:,.2f} kg/h',f'{ini[4]:,.2f}'],
        ['2小时储能','0 MWh  0 MW','0.00'],['4小时储能','0 MWh  0 MW','0.00'],
        ['合计','连续容量优化结果',f'{s["initial_investment_total_wanyuan"]:,.2f}']],[4.0,6.5,5.5])
    para(d,'储氢容量为可用氢气质量或标准体积，非罐体几何容积。设备容量尚未按整数台数及标准规格取整，投资口径未单列建安与配套工程。')
    heading(d,'三 年均成本与经济性')
    rows=[[n,f'{v:,.2f}',f'{v/s["cost_wanyuan_year"]*100:.2f}%'] for n,v in p.items()]
    rows.append(['合计',f'{s["cost_wanyuan_year"]:,.2f}','100.00%'])
    table(d,['成本项','万元每年','占比'],rows,[7,5,4])
    para(d,f'按年供氢3000万kg计算，模型单位成本为{s["unit_cost_yuan_kg"]:.2f}元/kg。该数值仅包含已建模的资产年化、运维、电量及需量成本，不是覆盖全部工程费用和运营费用的完整平准化制氢成本。')
    heading(d,'四 设备造价与财务边界')
    table(d,['项目','本方案单位投资','使用寿命'],[
        ['电解槽',f'{12000*factor:,.0f} 元/(Nm³/h)','20年'],
        ['电解槽等价交流功率口径',f'{12000*factor/((5+.1776)/.96):,.2f} 元/kW','额定初始100%负荷'],
        ['储氢罐',f'{3000*factor:,.0f} 元/kg  即 {3000*factor/11.1259:.2f} 元/Nm³','25年'],
        ['增压机',f'{30000*factor:,.0f} 元/(kg/h)','20年'],
        ['2小时储能','780 元/kWh','15年'],['4小时储能','700 元/kWh','15年']],[4.2,8.1,3.7])
    para(d,'投资按资产分别年化，回报率6%、所得税率25%、残值率5%，折旧年限取15年与设备寿命的较小值；年运维费统一按初始投资2%计入。储能功率投资系数为零，不代表PCS或接入工程已被确认免费。')
    heading(d,'五 电价与运行条件')
    table(d,['类别','每日区间','价格 元/kWh'],[
        ['谷段','11:00至18:00','0.11'],['平段','00:00至08:00  10:00至11:00\n18:00至19:00','0.28'],
        ['峰段','08:00至10:00  19:00至24:00','0.45']],[3,9,4])
    para(d,'使用2026年8760小时分时购电数据，各月采用相同区间，时刻为区间起点。不叠加原若羌节点电量杂费。峰段价格是内部负荷可避免的购电成本，不是上网销售收入。')
    para(d,'需量费保留原模型参数：基础30.4元/(kW·月)，附加项为0.0815×0.4785×730，合计约706.42元/(kW·年)。该口径尚未核验新项目适用性，不因设备造价变化而调整。')
    table(d,['参数','当前设置'],[
        ['年用氢及小时需求','30000吨  无检修  3424.66 kg/h'],['光伏及炼化负荷','均为零'],
        ['电解槽调节范围','额定产氢能力25%至120%'],['电解槽电耗','直流4.5467＋0.4533×负荷率  辅机0.1776 kWh/Nm³'],
        ['整流及衰减','整流效率96%  电堆电耗逐时按0.8%/年增长'],
        ['增压及储氢','增压2 kWh/kg  进出罐净流量不超过50%可用容量每小时'],
        ['储能效率及循环','充放各92%  年放电量不超过容量700倍'],
        ['电网及充放限制','只下网不上网  未强制仅平谷充峰段放']],[5,11])
    heading(d,'六 模型与求解口径')
    para(d,'采用连续线性规划，目标为购电费、需量费及各资产年化成本之和。逐时保持电力平衡与氢气物料平衡，储氢及储能均在年度首末闭环。电解槽电耗用切平面包络近似，不含整数台数、启停和爬坡。')
    para(d,'green为零时完全跳过绿电占比约束与缺口变量，允许正常购电补足储能损耗。储能所有小时含首小时均满足电量不超过容量；充放功率和不超过PCS容量，但并非严格整数互斥。Benders主问题选择容量，年度子问题优化运行并返回割平面。')
    heading(d,'七 储能结果解释与适用边界')
    if factor!=1:
        diff=base['cost_wanyuan_year']-s['cost_wanyuan_year']
        para(d,f'与原造价单年方案相比，年均成本减少{diff:,.2f}万元，降幅{diff/base["cost_wanyuan_year"]*100:.2f}%。原造价电解槽为{base["capacity"][2]:,.2f} Nm³/h、需量为{base["capacity"][1]/1000:.2f} MW；降价后电解槽容量增加、需量降低，但仍不配置电储能。')
    para(d,'不配置储能是在本方案电价、需量费、设备造价及供氢要求下的优化选择。储能需承担效率损耗及年化投资，新增充电需量可能带来较高费用；电解槽错峰生产与储氢罐亦可转移购电时段，构成电储能的替代路径。不能将单年零储能结论推广到其他电价、设备价格或供氢工况。')
    if factor==1:
        para(d,'原运行日志的固定其他容量增量分析显示，增加50 MWh的2小时或4小时储能，年度电费约节省42万元或40万元，低于相应年化投资约520万元或467万元。该测试未重新优化其余容量，不能作为全局临界造价。')
    else:
        para(d,'本方案定容已在第27轮达到精度条件，最佳配置来自第26轮。随后储能增量经济性复算长时间未返回，按用户要求终止。未生成原程序最终容量与逐年文件，因此本报告从最佳检查点与对应割平面恢复容量和成本，不引用未完成的差分收益或未保存的逐时调度。')
    para(d,'零容量边界的对偶值可能不唯一，不能直接外推为新增储能的实际收益；对偶折算临界造价仅用于诊断。判断储能经济性应以重新联合定容和可行调度比较为准。')
    para(d,'建安、土建、管线、升压接入、土地、设计及预备费未单列；单位投资是否已包含这些项目仍未确认。水耗、药剂、人工细分、堆体更换及营运资金等亦未独立建模。原参考报告储氢运维写1%，本次执行代码为2%，本报告以代码实际口径为准。')
    para(d,'本报告未进行十年后期电耗与检修校核，未生成新的全年逐时调度底稿或典型工况图。单年结果可用于初步比较，设备规格及长期供氢保障仍须进一步核验。')
    heading(d,'八 结果来源与保存状态')
    para(d,'参考文件为《轮台绿电制氢项目定容测算报告》，沿用其结果汇总、模型、边界及经济性分析的内容组织方式；原报告的光伏、炼化、原节点电价及十年工况不作为本报告结论。')
    para(d,f'源结果标识为 {s["source_tag"]}。本结果包包含原始日志与检查点、设备配置及投资CSV、年均成本CSV、测算结果汇总JSON、输入与代码快照及SHA256校验清单。代码快照记录归档时状态，实际算例参数以结果汇总和本报告为准。')
    para(d,'原程序全流程完成并保存容量与逐年结果。' if s['full_process_completed'] else '保存状态为主优化已收敛且附加后处理人为终止；恢复文件明确标识归档来源，不伪装为原程序完整输出。')
    path=dest/(subtitle+'定容测算报告.docx');d.save(path)
    return path


if __name__=='__main__':
    cases=[('01 原造价方案',1.0,SOURCE,'_bd1y_trk_offnone'),
           ('02 非储能设备降价40%方案',.6,SOURCE/'非储能设备造价倍率_0.6','_bd1y_trk_bx1_offnone')]
    saved=[archive(*c) for c in cases]
    for dest,s in saved:
        path=build(dest,s,saved[0][1]);print(path,flush=True)
        files=[]
        for p in sorted(dest.rglob('*')):
            if p.is_file() and p.name != '文件校验清单.json':files.append(dict(file=str(p.relative_to(dest)),bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest()))
        (dest/'文件校验清单.json').write_text(json.dumps(dict(archived_at=datetime.now().isoformat(),reference_sha256=hashlib.sha256(REF.read_bytes()).hexdigest(),files=files),ensure_ascii=False,indent=2),encoding='utf-8')
