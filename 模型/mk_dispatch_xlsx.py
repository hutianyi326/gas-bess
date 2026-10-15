# -*- coding: utf-8 -*-
"""附件：十年逐时调度底稿（仅调度数据）。

读 优化缓存/_dispatch_10y_trk_推荐.npz（经 dispatch_io 读取），每年一个工作表，8760 行。
输出：报告/附件 轮台绿电制氢项目十年逐时调度底稿.xlsx
用法：python mk_dispatch_xlsx.py
"""
import os
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from paths import DOC                                   # noqa: E402
DST = os.path.join(DOC, "附件 轮台绿电制氢项目十年逐时调度底稿.xlsx")
NPK = 11.1259
WEATHER = [2019, 2020, 2021, 2022, 2023] * 2

# (表头, 取数, 数字格式, 列宽)
COLS = [
    ("月", None, "0", 5), ("日", None, "0", 5), ("时（0～23）", None, "0", 7),
    ("光伏理论出力(MW)", lambda d: d["光伏可发"] / 1e3, "0.00", 12),
    ("光伏实际出力(MW)", lambda d: d["光伏利用"] / 1e3, "0.00", 12),
    ("弃光功率(MW)", lambda d: d["弃光"] / 1e3, "0.00", 11),
    ("电网购电功率(MW)", lambda d: d["下网"] / 1e3, "0.00", 12),
    ("储能充电功率(MW)", lambda d: d["储能充电"] / 1e3, "0.00", 12),
    ("储能放电功率(MW)", lambda d: d["储能放电"] / 1e3, "0.00", 12),
    ("储能剩余电量(MWh)", lambda d: d["储能电量"] / 1e3, "0.00", 12),
    ("电解槽用电功率(MW)", lambda d: d["电解槽"] / 1e3, "0.00", 12),
    ("增压机用电功率(MW)", lambda d: d["增压"] / 1e3, "0.00", 12),
    ("炼化装置用电负荷(MW)", lambda d: d["炼化"] / 1e3, "0.00", 13),
    ("电解槽制氢量(Nm³/h)", lambda d: d["产氢"], "0", 13),
    ("加氢装置用氢量(Nm³/h)", lambda d: d["取氢"], "0", 13),
    ("储氢罐储氢量(t)", lambda d: d["罐料位"] / NPK / 1e3, "0.00", 12),
]


def main():
    import sys
    sys.path.insert(0, HERE)
    import dispatch_io
    z = dispatch_io.load(report=True)
    cal = pd.date_range("2001-01-01", periods=8760, freq="h")     # 平年日历
    mo, dy, hr = cal.month.values, cal.day.values, cal.hour.values
    wb = Workbook(write_only=True)
    hfont = Font(name="宋体", size=10, bold=True)
    hfill = PatternFill("solid", fgColor="DDE4EC")
    thin = Side(style="thin", color="9AA5B1")
    hborder = Border(top=thin, bottom=thin, left=thin, right=thin)
    halign = Alignment(horizontal="center", vertical="center", wrap_text=True)
    nfont = Font(name="Arial", size=10)
    for yi in range(10):
        ws = wb.create_sheet("第%d年（%d年气象）" % (yi + 1, WEATHER[yi]))
        ws.freeze_panes = "D2"
        for i, (_, _, _, w) in enumerate(COLS):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.row_dimensions[1].height = 42
        head = []
        for name, *_ in COLS:
            c = WriteOnlyCell(ws, value=name)
            c.font, c.fill, c.border, c.alignment = hfont, hfill, hborder, halign
            head.append(c)
        ws.append(head)
        d = {k: v[yi] for k, v in z.items() if k not in ("x", "f", "epsy", "inv")}
        data = [None, None, None] + [f(d) for _, f, _, _ in COLS[3:]]
        fmts = [fm for _, _, fm, _ in COLS]
        for t in range(8760):
            row = []
            vals = [int(mo[t]), int(dy[t]), int(hr[t])] + [float(data[j][t]) for j in range(3, len(COLS))]
            for j, v in enumerate(vals):
                if j >= 3 and abs(v) < 5e-7:
                    v = 0.0
                c = WriteOnlyCell(ws, value=round(v, 4) if j >= 3 else v)
                c.font, c.number_format = nfont, fmts[j]
                row.append(c)
            ws.append(row)
        print("第%d年（天气 %d）写入 8760 行" % (yi + 1, WEATHER[yi]), flush=True)
    wb.save(DST)
    print("已存", DST, "%.1f MB" % (os.path.getsize(DST) / 1e6))


if __name__ == "__main__":
    main()
